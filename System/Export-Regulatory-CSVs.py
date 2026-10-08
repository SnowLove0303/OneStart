import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
import csv
import re
import shutil
from pathlib import Path
import openpyxl

# Paths
EXCEL_DIR_D = Path(r"D:\应用研究\统一启动器\法律法规物质限制清单_Excel导出")
EXCEL_DIR_F = Path(r"F:\APP Location\Guanzhi Tong\法律法规物质限制清单_Excel导出")

CSV_DIR_F = Path(r"F:\APP Location\Guanzhi Tong\法律法规物质限制清单_CSV导出")
CSV_DIR_D = Path(r"D:\应用研究\统一启动器\法律法规物质限制清单_CSV导出")

CSV_DIR_F.mkdir(parents=True, exist_ok=True)
CSV_DIR_D.mkdir(parents=True, exist_ok=True)

# Mapping of (workbook_stem, sheet_name) -> clean_csv_filename
SHEET_TO_CSV_NAME = {
    # 00
    ("00_法律法规判断器_六大物质限制清单汇总索引", "六大法规清单总览索引"): "00_法律法规判断器_六大物质限制清单汇总索引.csv",
    
    # 01 SVHC
    ("01_REACH_SVHC候选清单_253项", "物质全量清单(含组成员)"): "01_REACH_SVHC_01_物质全量清单(含组成员_545条).csv",
    ("01_REACH_SVHC候选清单_253项", "官方条目主清单(231项)"): "01_REACH_SVHC_02_官方条目主清单(231项).csv",
    ("01_REACH_SVHC候选清单_253项", "无明确标识条目(22项)"): "01_REACH_SVHC_03_无明确标识条目(22项_已补全CAS).csv",
    ("01_REACH_SVHC候选清单_253项", "版本与元数据"): "01_REACH_SVHC_04_版本与元数据.csv",

    # 02 Annex XVII
    ("02_REACH_附录XVII限制物质清单", "限制条目概览(86条)"): "02_REACH_附录XVII_01_限制条目概览(86条).csv",
    ("02_REACH_附录XVII限制物质清单", "受限物质明细与组成员(1868条)"): "02_REACH_附录XVII_02_受限物质明细与组成员(1868条).csv",
    ("02_REACH_附录XVII限制物质清单", "版本与元数据"): "02_REACH_附录XVII_03_版本与元数据.csv",

    # 03 RoHS
    ("03_EU_RoHS限制物质清单", "RoHS受限物质清单(160条)"): "03_EU_RoHS_01_受限物质清单(160条_含已补全类别CAS).csv",
    ("03_EU_RoHS限制物质清单", "版本与元数据"): "03_EU_RoHS_02_版本与元数据.csv",

    # 04 HSF-001
    ("04_HSF-001有害物质清单", "HSF-001有害物质清单(172条)"): "04_HSF-001_01_有害物质清单(172条_含已补全CAS).csv",
    ("04_HSF-001有害物质清单", "版本与元数据"): "04_HSF-001_02_版本与元数据.csv",

    # 05 BSBL
    ("05_BSBL物质限制清单_v8.0", "BSBL受限物质总表(1726条)"): "05_BSBL_01_受限物质总表(1726条_已修复日期并全量补全CAS).csv",
    ("05_BSBL物质限制清单_v8.0", "BSBL分类目录(51类)"): "05_BSBL_02_分类目录与条目数(51类).csv",
    ("05_BSBL物质限制清单_v8.0", "版本与元数据"): "05_BSBL_03_版本与元数据.csv",

    # 06 AfPS
    ("06_AfPS_GS_2019_01_PAK多环芳烃限制清单", "PAK限值清单(15单项+2合计)"): "06_AfPS_GS_2019_01_PAK_01_限值清单(15单项+2合计).csv",
    ("06_AfPS_GS_2019_01_PAK多环芳烃限制清单", "产品接触类别适用说明(5类)"): "06_AfPS_GS_2019_01_PAK_02_产品接触类别适用说明(5类).csv",
    ("06_AfPS_GS_2019_01_PAK多环芳烃限制清单", "法规来源与说明"): "06_AfPS_GS_2019_01_PAK_03_法规来源与说明.csv",

    # 07 Audit
    ("07_物质限制清单_缺失CAS补全明细对照表", "缺失CAS补全明细总表"): "07_物质限制清单_缺失CAS补全明细总表(442条).csv",
    ("07_物质限制清单_缺失CAS补全明细对照表", "补全统计分析"): "07_物质限制清单_各清单补全统计分析.csv"
}

def convert_excels_to_csv():
    print("=========================================================================")
    print("开始从法规限制清单 Excel 工作簿批量导出全量 CSV 文件 (UTF-8-SIG 编码)...")
    print("=========================================================================\n")
    
    # Check source Excel directory (prefer D:, fallback to F:)
    src_dir = EXCEL_DIR_D if EXCEL_DIR_D.exists() and any(EXCEL_DIR_D.glob("*.xlsx")) else EXCEL_DIR_F
    excel_files = sorted(src_dir.glob("*.xlsx"))
    
    if not excel_files:
        print(f"[ERROR] 未在目录找到 Excel 文件: {src_dir}")
        return []

    exported_files = []
    
    for ef in excel_files:
        if ef.name.startswith("~$") or "CAS补全版" in ef.name:
            continue
        stem = ef.stem
        wb = openpyxl.load_workbook(ef, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            csv_name = SHEET_TO_CSV_NAME.get((stem, sheet_name))
            if not csv_name:
                clean_sheet = re.sub(r'[\\/*?:"<>|]', '_', sheet_name)
                csv_name = f"{stem}_{clean_sheet}.csv"
                
            out_file_f = CSV_DIR_F / csv_name
            out_file_d = CSV_DIR_D / csv_name
            
            rows_written = 0
            with open(out_file_f, "w", newline="", encoding="utf-8-sig") as f_out:
                writer = csv.writer(f_out, quoting=csv.QUOTE_MINIMAL)
                for row in ws.iter_rows(values_only=True):
                    cleaned_row = ["" if v is None else str(v).strip() for v in row]
                    if not any(cleaned_row):
                        continue
                    writer.writerow(cleaned_row)
                    rows_written += 1
                    
            shutil.copy2(out_file_f, out_file_d)
            file_size = out_file_f.stat().st_size
            exported_files.append((csv_name, rows_written, file_size))
            print(f"[√] 成功导出 CSV: {csv_name:<55} ({rows_written:>5} 行, {file_size:>8} 字节)")

    print("\n=========================================================================")
    print(f"全套 CSV 清单批量导出完成！共计生成 {len(exported_files)} 个标准 CSV 文件。")
    print(f"主目录: {CSV_DIR_F}")
    print(f"镜像目录: {CSV_DIR_D}")
    print("=========================================================================")
    return exported_files

if __name__ == "__main__":
    convert_excels_to_csv()
